from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import pdfplumber
from PIL import Image, ImageChops, ImageStat
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "psyche-symbols-book"
SOURCE = ROOT / "outputs" / "PSYCHE-SYMBOLS-360.md"
DOCX = OUT / "The-Psyche-Symbols.docx"
PDF = OUT / "The-Psyche-Symbols.pdf"
RENDER_DIR = ROOT / "tmp" / "psyche-symbols-render"
QA = OUT / "book-final-qa.json"
POPPLER = Path(
    r"C:\Users\johnb\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe"
)


def expected_entries() -> list[dict]:
    pattern = re.compile(
        r"^(\d+)\. \*\*(Aries|Taurus|Gemini|Cancer|Leo|Virgo|Libra|Scorpio|Sagittarius|Capricorn|Aquarius|Pisces) (\d+) · ([^*]+)\.\*\* (.+)$"
    )
    entries = []
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        match = pattern.match(line)
        if match:
            entries.append(
                {
                    "index": int(match.group(1)),
                    "sign": match.group(2),
                    "degree": int(match.group(3)),
                    "title": match.group(4),
                    "image": match.group(5),
                }
            )
    return entries


def check_docx(entries: list[dict]) -> dict:
    doc = Document(DOCX)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    headings = [p.text.strip() for p in doc.paragraphs if p.style.name == "Heading 2"]
    expected_headings = [f"{e['sign']} {e['degree']}  {e['title']}" for e in entries]
    image_set = {e["image"] for e in entries}
    paragraph_set = set(paragraphs)
    labels = {label: 0 for label in ["INTERPRETATION", "SIGNAL", "SHADOW", "MIRROR", "PRACTICE"]}
    for text in paragraphs:
        for label in labels:
            if text.startswith(label + "  "):
                labels[label] += 1
    return {
        "heading_count": len(headings),
        "heading_order_ok": headings == expected_headings,
        "source_images_present": sum(1 for image in image_set if image in paragraph_set),
        "label_counts": labels,
        "title_metadata": doc.core_properties.title,
        "author_metadata": doc.core_properties.author,
        "sections": len(doc.sections),
    }


def render_pdf() -> list[Path]:
    if not POPPLER.exists():
        raise FileNotFoundError(POPPLER)
    RENDER_DIR.mkdir(parents=True, exist_ok=True)
    for old in RENDER_DIR.glob("page-*.png"):
        old.unlink()
    subprocess.run(
        [str(POPPLER), "-png", "-r", "90", str(PDF), str(RENDER_DIR / "page")],
        check=True,
        capture_output=True,
        text=True,
    )
    return sorted(RENDER_DIR.glob("page-*.png"))


def make_contact_sheets(images: list[Path]) -> list[Path]:
    sheets = []
    thumb_size = (216, 324)
    columns, rows = 4, 5
    chunk_size = columns * rows
    for old in RENDER_DIR.glob("contact-*.png"):
        old.unlink()
    for start in range(0, len(images), chunk_size):
        chunk = images[start : start + chunk_size]
        sheet = Image.new("RGB", (columns * thumb_size[0], rows * thumb_size[1]), "#c9c3b8")
        for idx, path in enumerate(chunk):
            with Image.open(path) as img:
                tile = img.convert("RGB")
                tile.thumbnail(thumb_size, Image.Resampling.LANCZOS)
                x = (idx % columns) * thumb_size[0] + (thumb_size[0] - tile.width) // 2
                y = (idx // columns) * thumb_size[1] + (thumb_size[1] - tile.height) // 2
                sheet.paste(tile, (x, y))
        page_start = start + 1
        page_end = start + len(chunk)
        output = RENDER_DIR / f"contact-{page_start:03d}-{page_end:03d}.png"
        sheet.save(output)
        sheets.append(output)
    return sheets


def image_checks(images: list[Path]) -> dict:
    sizes = []
    low_variance = []
    dark_edge_pages = []
    for page_number, path in enumerate(images, start=1):
        with Image.open(path) as img:
            rgb = img.convert("RGB")
            sizes.append(rgb.size)
            stat = ImageStat.Stat(rgb)
            variance = sum(stat.var) / 3
            if variance < 10:
                low_variance.append(page_number)
            width, height = rgb.size
            edge = Image.new("RGB", rgb.size, "white")
            edge.paste(rgb.crop((0, 0, width, 3)), (0, 0))
            edge.paste(rgb.crop((0, height - 3, width, height)), (0, height - 3))
            edge.paste(rgb.crop((0, 0, 3, height)), (0, 0))
            edge.paste(rgb.crop((width - 3, 0, width, height)), (width - 3, 0))
            diff = ImageChops.difference(edge, Image.new("RGB", rgb.size, "white"))
            if diff.getbbox() is not None:
                dark_edge_pages.append(page_number)
    return {
        "rendered_pages": len(images),
        "unique_render_sizes": sorted({f"{width}x{height}" for width, height in sizes}),
        "blank_or_low_variance_pages": low_variance,
        "ink_on_outermost_3px_pages": dark_edge_pages,
    }


def check_pdf(entries: list[dict]) -> dict:
    with pdfplumber.open(PDF) as pdf:
        page_texts = [(page.extract_text() or "") for page in pdf.pages]
        full_text = "\n".join(page_texts)
        sizes = sorted({f"{round(page.width, 2)}x{round(page.height, 2)}" for page in pdf.pages})
        # The illustrated edition intentionally uses an image-only cover page.
        # Treat a page with embedded raster art as nonblank even when it has no
        # extractable text; genuine blank pages have neither text nor images.
        blank_pages = [
            idx + 1
            for idx, (page, text) in enumerate(zip(pdf.pages, page_texts))
            if len(text.strip()) < 2 and not page.images
        ]
        edge_violations = []
        for page_idx, page in enumerate(pdf.pages, start=1):
            for word in page.extract_words():
                if word["x0"] < 18 or word["x1"] > page.width - 18 or word["top"] < 12 or word["bottom"] > page.height - 12:
                    edge_violations.append({"page": page_idx, "text": word["text"], "box": [word["x0"], word["top"], word["x1"], word["bottom"]]})
                    if len(edge_violations) >= 20:
                        break
            if len(edge_violations) >= 20:
                break
    expected_headings = [f"{e['sign']} {e['degree']} {e['title']}" for e in entries]
    normalized = re.sub(r"\s+", " ", full_text)
    found_headings = sum(1 for heading in expected_headings if heading in normalized)
    found_images = sum(1 for e in entries if e["image"] in normalized)
    return {
        "pages": len(page_texts),
        "page_sizes_points": sizes,
        "blank_pages": blank_pages,
        "symbol_headings_found": found_headings,
        "source_images_found": found_images,
        "interpretation_labels": len(re.findall(r"\bINTERPRETATION\b", full_text)),
        "signal_labels": len(re.findall(r"\bSIGNAL\b", full_text)),
        "shadow_labels": len(re.findall(r"\bSHADOW\b", full_text)),
        "mirror_labels": len(re.findall(r"\bMIRROR\b", full_text)),
        "practice_labels": len(re.findall(r"\bPRACTICE\b", full_text)),
        "edge_violations": edge_violations,
    }


def main() -> None:
    entries = expected_entries()
    docx = check_docx(entries)
    pdf = check_pdf(entries)
    images = render_pdf()
    image_qa = image_checks(images)
    sheets = make_contact_sheets(images)
    passed = (
        len(entries) == 360
        and docx["heading_count"] == 360
        and docx["heading_order_ok"]
        and docx["source_images_present"] == 360
        and all(count == 360 for count in docx["label_counts"].values())
        and pdf["page_sizes_points"] in (["432x648"], ["432.0x648.0"])
        and not pdf["blank_pages"]
        and pdf["symbol_headings_found"] == 360
        and pdf["source_images_found"] == 360
        and pdf["interpretation_labels"] == 360
        and pdf["signal_labels"] >= 360
        and pdf["shadow_labels"] >= 360
        and pdf["mirror_labels"] >= 360
        and pdf["practice_labels"] == 360
        and not pdf["edge_violations"]
        and image_qa["rendered_pages"] == pdf["pages"]
        and not image_qa["blank_or_low_variance_pages"]
        and not image_qa["ink_on_outermost_3px_pages"]
    )
    result = {
        "passed": passed,
        "entries": len(entries),
        "docx": docx,
        "pdf": pdf,
        "images": image_qa,
        "contact_sheets": [str(path) for path in sheets],
    }
    QA.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
    if not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
