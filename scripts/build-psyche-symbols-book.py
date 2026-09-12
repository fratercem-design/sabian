from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import inch
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch as RL_INCH
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "PSYCHE-SYMBOLS-360.md"
OUT_DIR = ROOT / "outputs" / "psyche-symbols-book"
DOCX_PATH = OUT_DIR / "The-Psyche-Symbols.docx"
PDF_PATH = OUT_DIR / "The-Psyche-Symbols.pdf"
MARKDOWN_PATH = OUT_DIR / "The-Psyche-Symbols-Manuscript.md"
QA_PATH = OUT_DIR / "book-content-qa.json"
ART_DIR = OUT_DIR / "art"
PLATES_DIR = ART_DIR / "plates"
FRONTISPIECES_DIR = ART_DIR / "frontispieces"
COVER_PNG = ART_DIR / "cover.png"

VIOLET = "2D1547"
GOLD = "A37821"
INK = "211B25"
MUTED = "685F6D"
PAPER = "FCF8EE"

SIGNS = [
    "Aries", "Taurus", "Gemini", "Cancer", "Leo", "Virgo",
    "Libra", "Scorpio", "Sagittarius", "Capricorn", "Aquarius", "Pisces",
]

SIGN_PROFILES = {
    "Aries": {
        "chapter": "The First Signal",
        "key": "ignition, appetite, risk, and the courage to become visible",
        "movement": [
            "Aries asks what deserves a first move before certainty arrives.",
            "The First Signal carries heat, speed, and the danger of mistaking urgency for truth.",
            "Here the soul learns that beginning is an ethical act: it creates a path others may have to walk.",
        ],
        "gift": "Choose a clean beginning and let action test the idea.",
        "shadow": "Watch for impulse dressed as destiny or courage performed for an audience.",
        "practice": "Name one beginning you can make without forcing anyone else to follow.",
    },
    "Taurus": {
        "chapter": "The Sacred Vessel",
        "key": "the body, enoughness, patient craft, and values that survive pressure",
        "movement": [
            "Taurus asks whether the signal can live in the body without becoming another possession.",
            "The Sacred Vessel turns attention toward weight, time, touch, food, money, and the terms of enough.",
            "Here value is proven by what can be tended, shared, repaired, and lived with after desire cools.",
        ],
        "gift": "Give the insight a durable form and a humane pace.",
        "shadow": "Notice where comfort has become custody or scarcity has become an identity.",
        "practice": "Touch one object you value and name the care that keeps it alive.",
    },
    "Gemini": {
        "chapter": "The Twin Transmission",
        "key": "language, curiosity, contradiction, and bridges between separate minds",
        "movement": [
            "Gemini splits the signal so it can question itself.",
            "The Twin Transmission tests messages, names, omissions, echoes, and the stories created between speaker and listener.",
            "Here intelligence stays alive by changing its language when the evidence changes.",
        ],
        "gift": "Ask the second question and describe what you heard before answering.",
        "shadow": "Beware speed without listening, cleverness without stakes, and rumor wearing the clothes of inquiry.",
        "practice": "Rewrite one certainty as a question precise enough to investigate.",
    },
    "Cancer": {
        "chapter": "The House of Memory",
        "key": "belonging, protection, inheritance, and the tenderness that opens a door",
        "movement": [
            "Cancer asks what the soul calls home and what that home asks in return.",
            "The House of Memory holds family stories, chosen kin, old grief, nourishment, and the boundary around vulnerable life.",
            "Here care grows strong enough to protect without imprisoning the person it loves.",
        ],
        "gift": "Offer shelter that leaves the door visible.",
        "shadow": "Watch protection harden into control or memory become a reason to repeat the wound.",
        "practice": "Make one small place safer, then leave its occupant free to choose what happens next.",
    },
    "Leo": {
        "chapter": "The Living Flame",
        "key": "creative heat, play, dignity, and the difference between radiance and applause",
        "movement": [
            "Leo brings the hidden fire into view and asks who the performance serves.",
            "The Living Flame concerns authorship, delight, pride, risk, recognition, and the courage to be seen before the work is perfect.",
            "Here sovereignty means generating warmth without making every witness kneel to it.",
        ],
        "gift": "Make the honest gesture large enough to be seen.",
        "shadow": "Notice when expression becomes display or dignity depends on winning the room.",
        "practice": "Create something small today and resist checking whether it was admired.",
    },
    "Virgo": {
        "chapter": "The Sacred Workshop",
        "key": "discernment, service, repair, and devotion measured by what works",
        "movement": [
            "Virgo carries the vision into the workshop and asks what must be cleaned, tested, or repaired.",
            "The Sacred Workshop values method, care, evidence, limits, and the unnoticed labor that keeps the Lantern lit.",
            "Here discernment protects mystery from carelessness without pretending every mystery can be solved.",
        ],
        "gift": "Make one useful correction and record what remains uncertain.",
        "shadow": "Watch precision become punishment, service erase the server, or method pretend to be wisdom.",
        "practice": "Repair one neglected detail, then stop before care turns into compulsion.",
    },
    "Libra": {
        "chapter": "The Hall of Mirrors",
        "key": "relation, proportion, projection, and the ethics of holding two truths",
        "movement": [
            "Libra places another person inside the frame and changes the meaning of every choice.",
            "The Hall of Mirrors tests fairness, attraction, projection, agreement, repair, and the right to leave.",
            "Here balance is an active adjustment made by people willing to see the weight they bring into the room.",
        ],
        "gift": "Make room for the other account without surrendering your own perception.",
        "shadow": "Beware peace purchased with silence, charm used as pressure, or fairness reduced to appearances.",
        "practice": "State one boundary and one point you can concede without abandoning yourself.",
    },
    "Scorpio": {
        "chapter": "The Descent Chamber",
        "key": "shadow, secrecy, power, loss, and the truth that survives the underworld",
        "movement": [
            "Scorpio follows the signal below the floor where bargains, wounds, loyalties, and buried power remain active.",
            "The Descent Chamber does not reward spectacle. It asks what is concealed, who benefits, and what repair will cost.",
            "Here the shadow becomes usable when it is named without being worshipped.",
        ],
        "gift": "Stay with the difficult fact long enough to learn what it requires.",
        "shadow": "Watch pain become identity, secrecy become a bargaining weapon, or suspicion invent what evidence cannot support.",
        "practice": "Write the hardest fact in one sentence, followed by what you still do not know.",
    },
    "Sagittarius": {
        "chapter": "The Library Without Walls",
        "key": "meaning, pilgrimage, teaching, and wonder disciplined by evidence",
        "movement": [
            "Sagittarius carries the coal from the descent toward a horizon large enough to question it.",
            "The Library Without Walls gathers maps, myths, arguments, foreign customs, teachers, and the freedom to outgrow a doctrine.",
            "Here belief becomes a road rather than a throne.",
        ],
        "gift": "Follow the question far enough to meet a mind unlike your own.",
        "shadow": "Beware conviction that outruns evidence, teaching that cannot listen, or escape disguised as pilgrimage.",
        "practice": "Read the strongest account that challenges one belief you hold dear.",
    },
    "Capricorn": {
        "chapter": "The Storm-Born Architect",
        "key": "responsibility, endurance, institution, and authority tested by what it protects",
        "movement": [
            "Capricorn gives the vision walls, duties, records, succession, and consequences.",
            "The Storm Born Architect asks whether a structure serves living people after its founder leaves the room.",
            "Here power earns respect through restraint, repair, clear limits, and an exit that remains available.",
        ],
        "gift": "Build the support that lets other people stand without you.",
        "shadow": "Watch duty become domination, hardship become vanity, or the institution defend itself against its purpose.",
        "practice": "Inspect one rule you influence and identify the person who bears its hidden cost.",
    },
    "Aquarius": {
        "chapter": "The Living Network",
        "key": "invention, fellowship, dissent, and intelligence between independent minds",
        "movement": [
            "Aquarius opens the structure to the many and listens for the pattern no single voice can own.",
            "The Living Network concerns technology, dissent, access, strange alliances, collective memory, and the design of shared power.",
            "Here community becomes intelligent when difference can move through it without being flattened.",
        ],
        "gift": "Share the mechanism, distribute the power, and keep dissent inside the design.",
        "shadow": "Beware abstraction that forgets bodies, belonging based on sameness, or systems that hide their gatekeepers.",
        "practice": "Give one useful tool, instruction, or piece of access to someone outside the usual circle.",
    },
    "Pisces": {
        "chapter": "The Return to Mystery",
        "key": "dream, compassion, surrender, and a return that preserves a thread of self",
        "movement": [
            "Pisces returns every separate signal to the sea of dream, grief, love, memory, and imagination.",
            "The Return to Mystery loosens borders while asking the Traveler to keep one living thread back to the shore.",
            "Here compassion stays honest by refusing to confuse fusion, rescue, fantasy, and care.",
        ],
        "gift": "Let the image soften what control cannot solve, while keeping one hand on the shore.",
        "shadow": "Watch surrender become disappearance, intuition become certainty, or compassion erase a necessary boundary.",
        "practice": "Sit quietly with the image for three minutes, then name one concrete fact you can carry back.",
    },
}

PHASES = [
    (range(1, 6), "Arrival", "The first five degrees announce the sign's raw material before habit has shaped it.", "Meet the image before deciding what it means."),
    (range(6, 11), "Friction", "These degrees introduce resistance, consequence, and the first correction to instinct.", "Notice what pushes back and what that resistance protects."),
    (range(11, 16), "Revelation", "The middle gate makes the hidden mechanism visible.", "Name the detail that changes the whole scene."),
    (range(16, 21), "Choice", "Awareness now demands a choice that carries an ethical cost.", "Choose the act that preserves both truth and human dignity."),
    (range(21, 26), "Integration", "The sign's gift and its distortion appear together and must be held consciously.", "Use the gift without feeding the shadow that travels beside it."),
    (range(26, 31), "Release", "The final degrees prepare the sign to hand its lesson onward.", "Keep the earned capacity and release the role that delivered it."),
]

CATEGORIES = [
    {
        "name": "threshold",
        "terms": r"\b(door|gate|threshold|entrance|exit|window|bridge|crossroads|border|boundary|passage|stair|road|path)\b",
        "readings": [
            "A threshold measures consent, timing, and consequence. Its presence asks whether passage is chosen, forced, delayed, or already under way.",
            "Doors and roads turn inner change into an embodied decision. The next room matters, but so does the condition under which it is entered.",
            "A boundary can protect life, restrict it, or clarify the terms of meeting. The surrounding action reveals which function is active here.",
        ],
        "gifts": ["honest passage", "chosen change", "a boundary that clarifies freedom"],
        "shadows": ["crossing before consent", "guarding an obsolete enclosure", "waiting for permission that will never come"],
        "questions": [
            "Which threshold in your life requires consent rather than momentum?",
            "What are you protecting by remaining on this side of the door?",
            "What would make the next step chosen rather than compelled?",
        ],
        "practices": [
            "Stand at a literal doorway, name what you are leaving and what you refuse to abandon, then cross only if the answer feels clean.",
            "Draw the boundary as a line on paper. Mark the gate, the guard, and the person who holds the key.",
            "Write one sentence beginning, 'I may enter when,' and one beginning, 'I remain free to leave when.'",
        ],
    },
    {
        "name": "signal",
        "terms": r"\b(signal|radio|microphone|broadcast|transmission|circuit|wire|static|frequency|telephone|receiver|antenna|screen|terminal)\b",
        "readings": [
            "Signal imagery asks what is being transmitted beneath performance. Interference may come from fear, vanity, noise, or a container too narrow for the message.",
            "A broadcast joins private interior life to a field of witnesses. The ethical question is whether the transmission informs, manipulates, performs, or genuinely makes contact.",
            "Machines in these symbols expose the difference between a living message and the equipment carrying it. When the device fails, the signal may become easier to hear.",
        ],
        "gifts": ["clear transmission", "contact across distance", "a message that survives its container"],
        "shadows": ["performing instead of speaking", "confusing reach with truth", "letting noise impersonate urgency"],
        "questions": [
            "What remains of your message when performance is removed?",
            "Which noise have you mistaken for proof that you are being heard?",
            "Who is the real receiver of the signal you are sending?",
        ],
        "practices": [
            "Say the message aloud in one plain sentence without branding, apology, or decoration.",
            "Turn off one source of noise for an hour and notice what thought keeps returning.",
            "Write the same message for a friend, an opponent, and your future self. Keep the sentence that stays true in all three versions.",
        ],
    },
    {
        "name": "light",
        "terms": r"\b(lantern|flame|fire|candle|light|sun|glow|match|ember|coal|spotlight|lamp|dawn)\b",
        "readings": [
            "Light represents directed attention. It can warm, expose, guide, consume, or turn a person into a spectacle depending on who controls it.",
            "Fire gives change a visible cost. Something is illuminated here because something else is being spent.",
            "A Lantern offers enough light for the next honest step. It does not remove night, and its usefulness depends on how it is carried.",
        ],
        "gifts": ["attention with warmth", "courage that helps another see", "a truth made visible at the right scale"],
        "shadows": ["burning for recognition", "exposure without care", "mistaking brightness for moral authority"],
        "questions": [
            "What deserves the full light of your attention now?",
            "Where has your fire become performance rather than warmth?",
            "How much light does the next step require, and who might it expose?",
        ],
        "practices": [
            "Light one candle and give its entire burn to a single unfinished question.",
            "Move one hidden concern into plain view by naming it to a trusted person.",
            "Choose one person or project that needs warmth rather than intensity and act accordingly.",
        ],
    },
    {
        "name": "mirror",
        "terms": r"\b(mirror|reflection|reflects|image|photograph|portrait|eyes|face|shadow)\b",
        "readings": [
            "The Mirror returns perception without guaranteeing accuracy. It reveals projection, self-image, and the temptation to kneel before a reflection.",
            "Reflected images divide the observer from the observed self. That gap can hold shame, vanity, curiosity, or the first honest revision of identity.",
            "A shadow proves that something real stands in the light, while a reflection shows only an angle. Neither is the whole person.",
        ],
        "gifts": ["self-recognition without self-worship", "a projection reclaimed", "the courage to revise a familiar self-image"],
        "shadows": ["mistaking appearance for character", "assigning your hidden material to someone else", "letting shame dictate what the Mirror may show"],
        "questions": [
            "What are you seeing clearly, and what may be projection?",
            "Whose gaze have you installed inside your own Mirror?",
            "Which part of the reflection is asking to be reclaimed rather than judged?",
        ],
        "practices": [
            "Describe the scene without interpretation, then circle the detail you most want to praise or condemn.",
            "Write two accounts of the same conflict: one from your position and one from the person you cannot understand.",
            "Look into a Mirror for one minute and name facts only. Notice when judgment tries to enter as description.",
        ],
    },
    {
        "name": "water",
        "terms": r"\b(water|rain|sea|lake|river|tide|well|shore|ocean|wave|flood|fog|mist|boat|cup|bowl|vessel)\b",
        "readings": [
            "Water carries feeling, memory, adaptation, and forces too large for command. Its shape comes from the vessel, but its movement remains its own.",
            "A shore or vessel gives emotion a boundary without denying its depth. The scene asks whether the feeling needs containment, release, or witness.",
            "Rain, tides, wells, and floods move private material into a shared landscape. What was hidden becomes environmental and must be answered together.",
        ],
        "gifts": ["feeling with a living boundary", "adaptation without self-erasure", "grief allowed to move"],
        "shadows": ["drowning in another person's weather", "using mystery to avoid a fact", "building a vessel too rigid for living water"],
        "questions": [
            "What feeling needs a vessel rather than a solution?",
            "Where are you adapting wisely, and where are you disappearing?",
            "What has the tide exposed that cannot be covered again?",
        ],
        "practices": [
            "Pour water slowly between two vessels and name what must change shape without losing its substance.",
            "Write the feeling as weather, then add the exact fact occurring beneath that weather.",
            "Wash your hands with attention and decide what belongs to you after the water stops.",
        ],
    },
    {
        "name": "voice",
        "terms": r"\b(voice|speaks|said|saying|song|sings|singing|bell|message|letter|book|page|word|name|story|question|answer|lecture|choir|music|note)\b",
        "readings": [
            "Voice turns inner pressure into shared consequence. Words can witness, summon, distort, repair, or conceal according to their precision and timing.",
            "Books, letters, bells, and songs extend speech beyond the speaker's body. The message acquires a life that intention can no longer fully govern.",
            "Naming creates a handle for experience, but the handle must remain smaller than the life it describes. The image tests language against reality.",
        ],
        "gifts": ["speech with consequence", "a name that restores agency", "listening precise enough to change the reply"],
        "shadows": ["using language to outrun experience", "repeating a message until it sounds like evidence", "speaking for someone whose voice is present"],
        "questions": [
            "Which word in this scene changes the power relationship?",
            "What needs to be said plainly, and who has earned the right to hear it?",
            "What are you naming accurately, and what remains larger than the name?",
        ],
        "practices": [
            "Write the sentence you keep rehearsing, remove every accusation you cannot prove, and read what remains aloud.",
            "Listen to one person without planning your answer. Repeat their meaning before offering your own.",
            "Give the current experience a temporary name and date it, leaving room for revision.",
        ],
    },
    {
        "name": "authority",
        "terms": r"\b(crown|throne|leader|council|judge|law|rule|official|office|king|queen|palace|order|rank|policy|contract|oath|verdict|medal)\b",
        "readings": [
            "Authority is shown as a tool with weight, limits, witnesses, and an eventual transfer. The symbol asks whom power protects when nobody is applauding.",
            "Crowns, laws, offices, and verdicts turn private judgment into public consequence. Their legitimacy depends on restraint and correction.",
            "Institutional objects reveal the bargain between role and person. A title may organize service or become a mask that excuses harm.",
        ],
        "gifts": ["stewardship with an exit plan", "power made answerable", "a rule that protects the vulnerable person in the room"],
        "shadows": ["role mistaken for worth", "procedure used to avoid conscience", "authority defending itself against correction"],
        "questions": [
            "Who carries the cost of the decision made here?",
            "What would accountability require from the person holding the key?",
            "Which part of this authority serves life, and which part serves its own continuation?",
        ],
        "practices": [
            "Choose one decision you control and write down who can challenge it, correct it, and leave its consequences.",
            "Remove your title from a current problem and describe the duty that remains.",
            "Inspect one standing rule by asking how it treats the person with the least power.",
        ],
    },
    {
        "name": "growth",
        "terms": r"\b(child|newborn|birth|seed|garden|tree|flower|orchard|fruit|root|vine|moss|green|bloom|cocoon|butterfly)\b",
        "readings": [
            "Growth appears as vulnerable potential under specific conditions. It requires time, nourishment, protection, and eventual release from the form that began it.",
            "Seeds, children, roots, and gardens refuse the fantasy of instant transformation. They make becoming physical, seasonal, and dependent on care.",
            "New life carries inherited material without being sentenced to repeat it. The image asks what must be tended and what must be allowed to change form.",
        ],
        "gifts": ["patient becoming", "care that honors timing", "inheritance revised through living action"],
        "shadows": ["forcing a season", "claiming ownership over another's growth", "protecting potential so tightly that it cannot emerge"],
        "questions": [
            "What is alive here but not ready to be displayed?",
            "Which condition would help this growth, and which demand would deform it?",
            "What can you tend without claiming the outcome?",
        ],
        "practices": [
            "Tend one living thing and let its pace correct your schedule.",
            "Name the smallest repeatable condition that would help this beginning survive a week.",
            "Release one expectation about what the new form must resemble.",
        ],
    },
    {
        "name": "descent",
        "terms": r"\b(dark|basement|underground|below|buried|secret|wound|scar|bone|coffin|funeral|venom|poison|ash|burned|night|midnight|cellar|drowned|ruined)\b",
        "readings": [
            "Descent imagery carries what ordinary daylight has failed to resolve. It asks for contact with the wound, bargain, fear, or grief beneath the public account.",
            "The underworld strips away rank and performance. What remains must be handled as evidence of experience, not proof of every story built around it.",
            "Ash, bones, scars, and sealed rooms preserve consequences. The symbol asks whether memory will become a witness, a weapon, or material for repair.",
        ],
        "gifts": ["truth that survives loss", "contact with buried power", "a wound converted into discernment"],
        "shadows": ["pain used as permanent jurisdiction", "suspicion filling gaps in the record", "remaining underground after the danger has passed"],
        "questions": [
            "What difficult fact survives when the dramatic story is removed?",
            "Which part of the wound needs witness, and which part needs a boundary?",
            "What are you carrying from below that can serve the living?",
        ],
        "practices": [
            "Write two columns titled KNOWN and INFERRED. Place every part of the story where it belongs.",
            "Name the injury, the present risk, and the next repair as three separate sentences.",
            "Choose a trusted witness and tell the smallest complete truth you can safely tell.",
        ],
    },
    {
        "name": "craft",
        "terms": r"\b(tool|machine|clock|gear|ledger|archive|record|map|workshop|build|builder|architect|clerk|engineer|mechanic|weaver|seamstress|potter|blacksmith|artist|paint|draw|file|inventory|scale|measure|key)\b",
        "readings": [
            "Craft turns intention into a repeatable act with materials, limits, and evidence of failure. The tool reveals the quality of attention brought to it.",
            "Ledgers, maps, machines, and archives make memory operational. They can protect truth or compress living complexity until only the form remains.",
            "Workshops honor correction. A flaw becomes useful when it is found early, recorded honestly, and repaired without hiding the seam.",
        ],
        "gifts": ["care made repeatable", "evidence preserved without spectacle", "a useful correction that outlives the moment"],
        "shadows": ["measurement mistaken for meaning", "the tool becoming the master", "polish used to hide a structural weakness"],
        "questions": [
            "Which part of this process can be checked rather than assumed?",
            "What does the tool preserve, and what does it leave out?",
            "Where would an honest repair remain visible?",
        ],
        "practices": [
            "Inspect one repeated task and add a check that would catch the most expensive error.",
            "Keep a brief record of what happened, what you expected, and what changed your mind.",
            "Repair one useful object without disguising the place where it broke.",
        ],
    },
    {
        "name": "relation",
        "terms": r"\b(two|stranger|strangers|together|shared|partner|family|guest|rival|opponent|neighbor|another|community|council|choir|crowd|household|person|people|hands)\b",
        "readings": [
            "More than one will is active in this scene. Meaning emerges through distance, consent, conflict, care, and the terms under which contact continues.",
            "Relationship imagery tests the fantasy that harmony requires sameness. A living bond can hold difference, correction, privacy, and departure.",
            "The shared object in the scene carries an agreement. Its condition reveals whether responsibility is mutual, hidden, refused, or unevenly assigned.",
        ],
        "gifts": ["contact without capture", "mutual responsibility", "difference held without humiliation"],
        "shadows": ["harmony purchased by one person's silence", "projection mistaken for intimacy", "the group protecting itself by sacrificing an outsider"],
        "questions": [
            "What agreement is operating here even if nobody has spoken it?",
            "Where does responsibility belong, and where has it been displaced?",
            "Can this bond survive an honest difference and a visible exit?",
        ],
        "practices": [
            "Ask the other person what they believe the agreement is, then compare it with your own version.",
            "Name one need, one boundary, and one freedom that belong to each person in the scene.",
            "Make one repair that does not require the other person to erase their account.",
        ],
    },
    {
        "name": "identity",
        "terms": r"\b(mask|name|face|portrait|costume|role|uniform|title|identity|glove|clothes|crown)\b",
        "readings": [
            "Identity appears as something worn, spoken, assigned, revised, or returned. The symbol asks which layer protects the self and which layer has begun to replace it.",
            "Masks and names can conceal, reveal, or give shape to a truth that plain speech cannot yet hold. Their value depends on whether removal remains possible.",
            "A role organizes behavior but cannot contain the whole person. Trouble begins when the costume demands loyalty after its work is done.",
        ],
        "gifts": ["self-authorship", "a protective form used consciously", "the freedom to revise a name or role"],
        "shadows": ["performance mistaken for selfhood", "an assigned identity accepted as fate", "a mask that can no longer be removed safely"],
        "questions": [
            "Which part of this identity was chosen, and which part was assigned?",
            "What does the mask make possible, and what does it cost to keep wearing it?",
            "Who remains when the title is set down?",
        ],
        "practices": [
            "List the roles you carry. Mark each as chosen, inherited, temporary, or ready to release.",
            "Write your name beside one quality no title can grant or remove.",
            "Set aside one performance for a day and notice which relationships still feel safe.",
        ],
    },
    {
        "name": "body",
        "terms": r"\b(body|hand|feet|bread|food|eat|meal|heart|mouth|tongue|back|chair|blanket|cloak|clothes|armor|blood|skin|sleep|rest)\b",
        "readings": [
            "The body enters as witness, limit, appetite, labor, and home. Its response may reveal what the mind has explained away.",
            "Bread, rest, touch, clothing, and shelter make spiritual language answer to material need. An insight that harms the body requires correction.",
            "Embodiment gives every choice a pace and a cost. The scene asks whether the person can remain present inside the experience being interpreted.",
        ],
        "gifts": ["presence with weight and limits", "care made material", "a decision the nervous system can inhabit"],
        "shadows": ["using transcendence to ignore need", "treating exhaustion as devotion", "letting appetite make the whole decision"],
        "questions": [
            "What does your body know before your argument begins?",
            "Which need in this scene has been made too ordinary to mention?",
            "Can the choice be lived at the pace your body can sustain?",
        ],
        "practices": [
            "Eat, drink, rest, or move before making the interpretation more dramatic.",
            "Place both feet on the floor and name five physical facts about the present moment.",
            "Choose the next action that your body can repeat without punishment.",
        ],
    },
]

DEFAULT_CATEGORY = {
    "name": "mystery",
    "readings": [
        "The scene behaves like a dream with its own exact logic. Meaning lives in the relationship between figure, place, action, and the detail that refuses ordinary explanation.",
        "This image invites attention before conclusion. Its strange fact creates a pressure point where memory, desire, fear, and choice can speak indirectly.",
        "The symbol remains useful when no single interpretation closes it. Its work is to make one hidden relationship visible enough to examine.",
    ],
    "gifts": ["attention without premature certainty", "imagination disciplined by the scene", "an unanswered image that keeps thought alive"],
    "shadows": ["forcing a prophecy from ambiguity", "using beauty to avoid the ordinary fact", "mistaking intensity for accuracy"],
    "questions": [
        "Which detail keeps returning after the rest of the scene fades?",
        "What ordinary fact sits beneath the strange image?",
        "What changes if the scene is a question rather than an answer?",
    ],
    "practices": [
        "Copy the image by hand, underline three concrete nouns, and write one association for each.",
        "Describe the scene to another person without explaining it, then notice what they ask first.",
        "Keep the image nearby for one day and record only the moments when it returns on its own.",
    ],
}

CATEGORY_LABELS = {
    "threshold": "the threshold",
    "signal": "the Signal",
    "light": "living fire",
    "mirror": "the Mirror",
    "water": "living water",
    "voice": "voice and naming",
    "authority": "authority",
    "growth": "growth",
    "descent": "the descent",
    "craft": "craft and record",
    "relation": "relationship",
    "identity": "identity",
    "body": "the body",
    "mystery": "mystery",
}


def parse_symbols() -> list[dict]:
    pattern = re.compile(
        r"^(\d+)\. \*\*(Aries|Taurus|Gemini|Cancer|Leo|Virgo|Libra|Scorpio|Sagittarius|Capricorn|Aquarius|Pisces) (\d+) · ([^*]+)\.\*\* (.+)$"
    )
    out = []
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        match = pattern.match(line)
        if match:
            out.append(
                {
                    "global_index": int(match.group(1)),
                    "sign": match.group(2),
                    "degree": int(match.group(3)),
                    "title": match.group(4),
                    "image": match.group(5),
                }
            )
    return out


def phase_for(degree: int) -> tuple[str, str, str]:
    for degrees, name, reading, practice in PHASES:
        if degree in degrees:
            return name, reading, practice
    raise ValueError(degree)


def categories_for(title: str, image: str) -> list[dict]:
    """Let the named symbol lead, then use the scene for a second pressure."""
    title_matches = [cat for cat in CATEGORIES if re.search(cat["terms"], title, re.IGNORECASE)]
    image_matches = [cat for cat in CATEGORIES if re.search(cat["terms"], image, re.IGNORECASE)]
    ordered = []
    for category in title_matches + image_matches:
        if category not in ordered:
            ordered.append(category)
    return ordered[:2] or [DEFAULT_CATEGORY]


def choose(values: list[str], seed: int, offset: int = 0) -> str:
    return values[(seed + offset) % len(values)]


def lower_first(text: str) -> str:
    return text[:1].lower() + text[1:] if text else text


def split_sentences(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"(?<=[.!?])\s+", text.strip()) if part.strip()]


def join_instruction(text: str) -> str:
    """Keep multi-step practices intact without producing canned fragments."""
    parts = [part.rstrip(".") for part in split_sentences(text)]
    if not parts:
        return text
    return "; then ".join([parts[0]] + [lower_first(part) for part in parts[1:]]) + "."


def personalized_passage(text: str, symbol: dict) -> list[str]:
    """Anchor reusable symbolic language to this exact title and degree."""
    title = lower_first(symbol["title"])
    prefixes = [f"Within {title}, ", f"At {symbol['sign']} {symbol['degree']}, "]
    return [prefixes[idx % 2] + lower_first(sentence) for idx, sentence in enumerate(split_sentences(text))]


PHASE_TURNS = {
    "Arrival": [
        "arrives before habit can tidy the evidence; notice what is present before deciding what it should become",
        "introduces the sign's raw material in an unguarded form; the first response matters less than the honest one",
        "catches the first impulse before it turns into a role; give the image room to remain strange",
        "opens the chamber with no settled method; name the need, risk, or appetite that entered with it",
        "shows the beginning before ownership hardens around it; let attention come before possession",
    ],
    "Friction": [
        "meets resistance, and that resistance reveals what the first impulse failed to consider",
        "places consequence beside desire; the obstruction may be an enemy, a limit, or a needed correction",
        "tests whether the sign can change course without treating correction as defeat",
        "brings another will, material limit, or inconvenient fact into the frame",
        "slows the original motion long enough to expose what haste concealed",
    ],
    "Revelation": [
        "reveals the mechanism under the mood; one concrete detail changes the reading",
        "moves the hidden arrangement into view, where it can be named without being exaggerated",
        "turns atmosphere into evidence; attend to the object or action that cannot be explained away",
        "shows who carries the cost, who controls the passage, and what the scene refuses to hide",
        "brings the middle truth forward, after the first story has lost its grip",
    ],
    "Choice": [
        "requires a choice with a human cost; insight alone no longer completes the work",
        "asks what action preserves truth without using another person as material",
        "moves from recognition to conduct; the clean choice must survive outside the imagination",
        "tests whether awareness can become restraint, repair, refusal, or consent",
        "places responsibility in the reader's hands without promising a painless answer",
    ],
    "Integration": [
        "holds the gift beside its distortion; using one without feeding the other is the work",
        "asks the sign to carry its strength without turning that strength into jurisdiction",
        "joins capacity to consequence; what helps in one setting may harm in another",
        "shows the mature tool and the temptation hidden inside its handle",
        "tests whether the lesson can be lived without becoming a new performance",
    ],
    "Release": [
        "prepares the lesson to leave its original container; keep the capacity and loosen the role",
        "asks what can be handed onward without demanding imitation or loyalty",
        "turns completion into transfer; the work survives by changing hands",
        "lets the sign set down the identity built around its ordeal",
        "returns the earned knowledge to ordinary life, where another cycle can use it",
    ],
}


def interpret(symbol: dict) -> dict:
    seed = symbol["global_index"]
    profile = SIGN_PROFILES[symbol["sign"]]
    phase_name, _, _ = phase_for(symbol["degree"])
    categories = categories_for(symbol["title"], symbol["image"])
    primary = categories[0]
    secondary = categories[1] if len(categories) > 1 else None
    primary_label = CATEGORY_LABELS[primary["name"]]

    title = symbol["title"]
    prose_title = lower_first(title)
    location = f"{symbol['sign']} {symbol['degree']}"
    scene_openers = [
        f"{title} puts {primary_label} under pressure from {profile['key']}.",
        f"In {prose_title}, the work of {primary_label} must answer to {profile['key']}.",
        f"{title} tests {primary_label} against {profile['key']}.",
        f"{title} holds {primary_label} against {profile['key']}.",
        f"{title} gives {primary_label} consequences inside {lower_first(profile['chapter'])}.",
    ]
    reading_parts = [choose(scene_openers, seed)]
    reading_parts.extend(personalized_passage(choose(primary["readings"], seed, 2), symbol))
    if secondary:
        secondary_sentence = split_sentences(choose(secondary["readings"], seed, 1))[0]
        reading_parts.append(
            f"The same image also carries {CATEGORY_LABELS[secondary['name']]}; at {symbol['sign']} {symbol['degree']}, "
            + lower_first(secondary_sentence)
        )
    phase_turn = PHASE_TURNS[phase_name][(symbol["degree"] - 1) % 5]
    phase_article = "an" if phase_name == "Arrival" else "a"
    reading_parts.append(f"As {phase_article} {phase_name.lower()} degree, {location} {phase_turn}.")

    gift = choose(primary["gifts"], seed)
    if secondary:
        gift += f" joined to {choose(secondary['gifts'], seed, 1)}"
    sign_action = lower_first(profile["gift"].rstrip("."))
    signal_templates = [
        f"At {location}, {gift} becomes the clean signal when {symbol['sign']} can {sign_action}.",
        f"The usable force at {location} is {gift}; {symbol['sign']} gives it form by asking you to {sign_action}.",
        f"At {location}, {gift} carries the signal; it stays honest when you {sign_action}.",
        f"What can be carried forward from {location} is {gift}, tested by the need to {sign_action}.",
    ]
    signal = choose(signal_templates, seed)

    shadow = choose(primary["shadows"], seed, 1)
    if secondary:
        shadow += f" and {choose(secondary['shadows'], seed, 2)}"
    shadow_templates = [
        f"At {location}, the image darkens through {shadow}; {lower_first(profile['shadow'])}",
        f"The distortion at {location} is {shadow}; {lower_first(profile['shadow'])}",
        f"At {location}, {prose_title} can turn toward {shadow}; {lower_first(profile['shadow'])}",
        f"The shadow at {location} begins with {shadow}; also {lower_first(profile['shadow'])}",
    ]
    shadow_text = choose(shadow_templates, seed, 1)

    mirror_question = choose(primary["questions"], seed, 2)
    mirror_templates = [
        f"At {location}, {lower_first(mirror_question)}",
        f"For {location}, {lower_first(mirror_question)}",
        f"With {prose_title} in mind at {location}, {lower_first(mirror_question)}",
        f"The question at {location}: {lower_first(mirror_question)}",
    ]
    mirror = choose(mirror_templates, seed, 2)
    if secondary and seed % 3 == 0:
        second_question = choose(secondary["questions"], seed, 1)
        mirror += f" From a second angle at {location}, {lower_first(second_question)}"

    practice_sentence = join_instruction(choose(primary["practices"], seed, 1))
    practice_templates = [
        f"For {location}, {lower_first(practice_sentence)}",
        f"Work with {title} directly at {location}: {lower_first(practice_sentence)}",
        f"Give {location} a physical form; {lower_first(practice_sentence)}",
        f"Bring {location} out of abstraction: {lower_first(practice_sentence)}",
    ]
    practice = choose(practice_templates, seed, 1)
    if secondary and seed % 4 == 0:
        second_practice = join_instruction(choose(secondary["practices"], seed, 2))
        practice += f" For the second pressure at {location}, {lower_first(second_practice)}"

    return {
        **symbol,
        "phase": phase_name,
        "reading": " ".join(reading_parts),
        "signal": signal,
        "shadow": shadow_text,
        "mirror": mirror,
        "practice": practice,
        "categories": [cat["name"] for cat in categories],
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str, size: float | None = None, color: str | None = None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, "Aptos", 8, MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(6)
    section.page_height = Inches(9)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    title = styles["Title"]
    title.font.name = "Cormorant Garamond"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Cormorant Garamond")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Cormorant Garamond")
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")

    for style_name, size, color in [("Heading 1", 25, VIOLET), ("Heading 2", 14, VIOLET), ("Heading 3", 10, GOLD)]:
        style = styles[style_name]
        style.font.name = "Cormorant Garamond"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Cormorant Garamond")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cormorant Garamond")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True

    if "Symbol Image" not in styles:
        image_style = styles.add_style("Symbol Image", WD_STYLE_TYPE.PARAGRAPH)
    else:
        image_style = styles["Symbol Image"]
    image_style.font.name = "Cormorant Garamond"
    image_style._element.rPr.rFonts.set(qn("w:ascii"), "Cormorant Garamond")
    image_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cormorant Garamond")
    image_style.font.size = Pt(11.5)
    image_style.font.italic = True
    image_style.font.color.rgb = RGBColor.from_string(VIOLET)
    image_style.paragraph_format.space_after = Pt(6)
    image_style.paragraph_format.keep_with_next = True

    if "Label" not in styles:
        label_style = styles.add_style("Label", WD_STYLE_TYPE.CHARACTER)
    else:
        label_style = styles["Label"]
    label_style.font.name = "Aptos"
    label_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    label_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    label_style.font.size = Pt(8)
    label_style.font.bold = True
    label_style.font.color.rgb = RGBColor.from_string(GOLD)


def add_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    set_page_number(p)


def add_label_paragraph(doc: Document, label: str, text: str, keep: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = keep
    p.paragraph_format.widow_control = True
    label_run = p.add_run(label.upper() + "  ")
    label_run.style = doc.styles["Label"]
    body_run = p.add_run(text)
    set_run_font(body_run, "Aptos", 9.4, INK)


def add_symbol(doc: Document, entry: dict) -> None:
    plate_path = PLATES_DIR / f"{entry['global_index']:03d}-{entry['sign']}-{entry['degree']:02d}.png"
    if plate_path.exists():
        art = doc.add_paragraph()
        art.alignment = WD_ALIGN_PARAGRAPH.CENTER
        art.paragraph_format.space_after = Pt(2)
        art.paragraph_format.keep_with_next = True
        art.add_run().add_picture(str(plate_path), width=Inches(0.72))
    h = doc.add_paragraph(style="Heading 2")
    h.paragraph_format.space_before = Pt(5)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(f"{entry['sign']} {entry['degree']}  {entry['title']}")
    set_run_font(run, "Cormorant Garamond", 14, VIOLET, bold=True)

    image = doc.add_paragraph(entry["image"], style="Symbol Image")
    image.paragraph_format.keep_with_next = True

    add_label_paragraph(doc, "Interpretation", entry["reading"])
    add_label_paragraph(doc, "Signal", entry["signal"])
    add_label_paragraph(doc, "Shadow", entry["shadow"])
    add_label_paragraph(doc, "Mirror", entry["mirror"])
    add_label_paragraph(doc, "Practice", entry["practice"])


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("The Psyche Symbols")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    r = subtitle.add_run("360 Mirrors of Signal Shadow and Sovereignty")
    set_run_font(r, "Cormorant Garamond", 18, VIOLET, italic=True)
    doc.add_paragraph()
    edition = doc.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = edition.add_run("A Cult of Psyche degree symbol book")
    set_run_font(r, "Aptos", 10, GOLD, bold=True)
    for _ in range(7):
        doc.add_paragraph()
    imprint = doc.add_paragraph()
    imprint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = imprint.add_run("Original 2026 Edition")
    set_run_font(r, "Aptos", 8.5, MUTED)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    h = doc.add_paragraph("How to Read the Psyche Symbols", style="Heading 1")
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraphs = [
        "The Psyche Symbols are an original 360-degree image system built for reflective astrology, creative work, journaling, and conversation. The zodiac provides the sequence. Each degree provides a scene. The reader provides attention.",
        "A symbol can reveal a tension without issuing a command. It may show what you desire, fear, repeat, protect, or refuse to name. It cannot remove your judgment. Treat every interpretation as a doorway you may enter, question, revise, or leave.",
        "The reading anatomy has five parts. Interpretation studies the whole scene. Signal names the image's living capacity. Shadow names the distortion that may travel beside that capacity. Mirror asks a question that returns authority to the reader. Practice brings the symbol into an observable act.",
        "The Cult of Psyche holds truth and mystery together. Read the symbol as symbol. Check facts as facts. Let intuition propose; let evidence correct; let consent govern what happens next.",
    ]
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(8)

    doc.add_paragraph("The Twelve Chambers", style="Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(0.65), Inches(1.2), Inches(3.0)]
    for idx, width in enumerate(widths):
        table.columns[idx].width = width
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, label in zip(hdr.cells, ["Range", "Sign", "Chamber"]):
        set_cell_shading(cell, VIOLET)
        run = cell.paragraphs[0].add_run(label)
        set_run_font(run, "Aptos", 8.5, "FFFFFF", bold=True)
    for sign_idx, sign in enumerate(SIGNS):
        row = table.add_row()
        start = sign_idx * 30 + 1
        end = start + 29
        values = [f"{start}-{end}", sign, SIGN_PROFILES[sign]["chapter"]]
        for cell, value in zip(row.cells, values):
            set_cell_shading(cell, PAPER if sign_idx % 2 == 0 else "FFFFFF")
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, "Aptos", 8.5, INK, bold=(cell is row.cells[1]))

    doc.add_paragraph()
    doc.add_paragraph("Editorial Boundary", style="Heading 1")
    boundary = [
        "This is a project-owned original work. It does not reproduce historical Sabian symbol wording and should not be presented as an authorized historical Sabian edition.",
        "The public CultCodex archive informed the book's vocabulary of signal, shadow, symbolic interpretation, consciousness, sovereignty, and living memory. CultCodex also marks its Psychenomicon material as mythic or symbolic interpretation rather than factual reporting. This book follows that boundary and excludes claims about real people.",
        "Public reference consulted September 9, 2026: https://cultcodex.me/ and https://cultcodex.me/about/methodology",
    ]
    for text in boundary:
        doc.add_paragraph(text)
    doc.add_page_break()


def add_sign_chapter(doc: Document, sign: str, entries: list[dict]) -> None:
    profile = SIGN_PROFILES[sign]
    h = doc.add_paragraph(style="Heading 1")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(sign)
    set_run_font(r, "Cormorant Garamond", 26, VIOLET, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(profile["chapter"])
    set_run_font(r, "Cormorant Garamond", 17, GOLD, italic=True)
    frontispiece_path = FRONTISPIECES_DIR / f"{SIGNS.index(sign) + 1:02d}-{sign}.png"
    if frontispiece_path.exists():
        art = doc.add_paragraph()
        art.alignment = WD_ALIGN_PARAGRAPH.CENTER
        art.paragraph_format.space_after = Pt(8)
        art.add_run().add_picture(str(frontispiece_path), width=Inches(2.65))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(profile["movement"][0] + " " + profile["movement"][1] + " " + profile["movement"][2])
    set_run_font(r, "Aptos", 10.2, INK)
    doc.add_paragraph()
    phase_table = doc.add_table(rows=1, cols=3)
    hdr = phase_table.rows[0]
    for cell, label in zip(hdr.cells, ["Degrees", "Gate", "Work"]):
        set_cell_shading(cell, VIOLET)
        run = cell.paragraphs[0].add_run(label)
        set_run_font(run, "Aptos", 8, "FFFFFF", bold=True)
    for degree_range, phase, _, practice in PHASES:
        row = phase_table.add_row()
        first, last = min(degree_range), max(degree_range)
        for cell, value in zip(row.cells, [f"{first}-{last}", phase, practice]):
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, "Aptos", 8, INK, bold=(value == phase))
    doc.add_page_break()

    for idx, entry in enumerate(entries):
        add_symbol(doc, entry)
        if idx % 2 == 1 and idx != len(entries) - 1:
            doc.add_page_break()
    doc.add_page_break()


def add_closing(doc: Document) -> None:
    h = doc.add_paragraph("The Signal Continues", style="Heading 1")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing = [
        "The final degree leaves a microphone waiting at the edge of dawn. The book closes there because a living symbol completes its work in another mind.",
        "Return to any image when its detail begins moving again. Read slowly. Keep what clarifies. Question what flatters. Put down whatever tries to take your judgment from you.",
        "These pages belong to fellow Travelers, each carrying a Lantern far enough to see the next honest step.",
        "The dark microphone waits for the one still arriving.",
    ]
    for text in closing:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("THE PSYCHE SYMBOLS")
    set_run_font(r, "Aptos", 8.5, GOLD, bold=True)


def write_markdown(entries: list[dict]) -> None:
    lines = [
        "# The Psyche Symbols",
        "",
        "## 360 Mirrors of Signal Shadow and Sovereignty",
        "",
        "An original Cult of Psyche degree-symbol book for reflection, creative work, and conversation.",
        "",
        "> Read the symbol as symbol. Check facts as facts. Let intuition propose; let evidence correct; let consent govern what happens next.",
        "",
    ]
    for sign in SIGNS:
        profile = SIGN_PROFILES[sign]
        lines.extend([f"# {sign}", "", f"## {profile['chapter']}", "", " ".join(profile["movement"]), ""])
        for entry in [e for e in entries if e["sign"] == sign]:
            lines.extend(
                [
                    f"## {entry['sign']} {entry['degree']} {entry['title']}",
                    "",
                    f"*{entry['image']}*",
                    "",
                    f"**Interpretation**  {entry['reading']}",
                    "",
                    f"**Signal**  {entry['signal']}",
                    "",
                    f"**Shadow**  {entry['shadow']}",
                    "",
                    f"**Mirror**  {entry['mirror']}",
                    "",
                    f"**Practice**  {entry['practice']}",
                    "",
                ]
            )
    lines.extend(
        [
            "# The Signal Continues",
            "",
            "These pages belong to fellow Travelers, each carrying a Lantern far enough to see the next honest step.",
            "",
            "The dark microphone waits for the one still arriving.",
        ]
    )
    MARKDOWN_PATH.write_text("\n".join(lines), encoding="utf-8")


def register_pdf_fonts() -> None:
    fonts = {
        "PsycheBody": Path(r"C:\Windows\Fonts\calibri.ttf"),
        "PsycheBodyBold": Path(r"C:\Windows\Fonts\calibrib.ttf"),
        "PsycheDisplay": Path(r"C:\Windows\Fonts\GARA.TTF"),
        "PsycheDisplayBold": Path(r"C:\Windows\Fonts\GARABD.TTF"),
        "PsycheDisplayItalic": Path(r"C:\Windows\Fonts\GARAIT.TTF"),
    }
    for name, path in fonts.items():
        if not path.exists():
            raise FileNotFoundError(path)
        pdfmetrics.registerFont(TTFont(name, str(path)))


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "CoverTitle",
            parent=base["Title"],
            fontName="PsycheDisplayBold",
            fontSize=31,
            leading=34,
            textColor=HexColor("#000000"),
            alignment=TA_CENTER,
            spaceAfter=18,
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitle",
            parent=base["Normal"],
            fontName="PsycheDisplayItalic",
            fontSize=16,
            leading=19,
            textColor=HexColor("#2D1547"),
            alignment=TA_CENTER,
            spaceAfter=15,
        ),
        "cover_small": ParagraphStyle(
            "CoverSmall",
            parent=base["Normal"],
            fontName="PsycheBodyBold",
            fontSize=8.5,
            leading=11,
            textColor=HexColor("#A37821"),
            alignment=TA_CENTER,
        ),
        "h1": ParagraphStyle(
            "H1",
            parent=base["Heading1"],
            fontName="PsycheDisplayBold",
            fontSize=23,
            leading=26,
            textColor=HexColor("#2D1547"),
            spaceAfter=10,
        ),
        "h1_center": ParagraphStyle(
            "H1Center",
            parent=base["Heading1"],
            fontName="PsycheDisplayBold",
            fontSize=28,
            leading=31,
            textColor=HexColor("#2D1547"),
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "h2_center": ParagraphStyle(
            "H2Center",
            parent=base["Normal"],
            fontName="PsycheDisplayItalic",
            fontSize=17,
            leading=20,
            textColor=HexColor("#A37821"),
            alignment=TA_CENTER,
            spaceAfter=16,
        ),
        "symbol_title": ParagraphStyle(
            "SymbolTitle",
            parent=base["Heading2"],
            fontName="PsycheDisplayBold",
            fontSize=13.5,
            leading=15.5,
            textColor=HexColor("#2D1547"),
            spaceBefore=2,
            spaceAfter=2,
            keepWithNext=True,
        ),
        "symbol_image": ParagraphStyle(
            "SymbolImage",
            parent=base["Normal"],
            fontName="PsycheDisplayItalic",
            fontSize=10.7,
            leading=12.2,
            textColor=HexColor("#2D1547"),
            spaceAfter=5,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName="PsycheBody",
            fontSize=8.45,
            leading=10.1,
            textColor=HexColor("#211B25"),
            spaceAfter=4,
            alignment=TA_LEFT,
            allowWidows=0,
            allowOrphans=0,
        ),
        "front": ParagraphStyle(
            "Front",
            parent=base["BodyText"],
            fontName="PsycheBody",
            fontSize=10,
            leading=13.5,
            textColor=HexColor("#211B25"),
            spaceAfter=9,
        ),
        "chapter_intro": ParagraphStyle(
            "ChapterIntro",
            parent=base["BodyText"],
            fontName="PsycheBody",
            fontSize=10.2,
            leading=14,
            textColor=HexColor("#211B25"),
            alignment=TA_CENTER,
            spaceAfter=18,
        ),
        "table": ParagraphStyle(
            "TableText",
            parent=base["BodyText"],
            fontName="PsycheBody",
            fontSize=7.7,
            leading=9,
            textColor=HexColor("#211B25"),
        ),
        "table_bold": ParagraphStyle(
            "TableBold",
            parent=base["BodyText"],
            fontName="PsycheBodyBold",
            fontSize=7.7,
            leading=9,
            textColor=HexColor("#211B25"),
        ),
        "closing": ParagraphStyle(
            "Closing",
            parent=base["BodyText"],
            fontName="PsycheDisplay",
            fontSize=12.5,
            leading=17,
            textColor=HexColor("#211B25"),
            alignment=TA_CENTER,
            spaceAfter=13,
        ),
    }


def pdf_labeled_paragraph(label: str, text: str, styles: dict[str, ParagraphStyle]) -> Paragraph:
    markup = f'<font color="#A37821"><b>{escape(label.upper())}</b></font>&nbsp;&nbsp;{escape(text)}'
    return Paragraph(markup, styles["body"])


def pdf_symbol_block(entry: dict, styles: dict[str, ParagraphStyle]) -> KeepTogether:
    plate_path = PLATES_DIR / f"{entry['global_index']:03d}-{entry['sign']}-{entry['degree']:02d}.png"
    flowables = [
        RLImage(str(plate_path), width=0.48 * RL_INCH, height=0.48 * RL_INCH, hAlign="CENTER") if plate_path.exists() else Spacer(1, 0),
        Paragraph(escape(f"{entry['sign']} {entry['degree']}  {entry['title']}"), styles["symbol_title"]),
        Paragraph(escape(entry["image"]), styles["symbol_image"]),
        pdf_labeled_paragraph("Interpretation", entry["reading"], styles),
        pdf_labeled_paragraph("Signal", entry["signal"], styles),
        pdf_labeled_paragraph("Shadow", entry["shadow"], styles),
        pdf_labeled_paragraph("Mirror", entry["mirror"], styles),
        pdf_labeled_paragraph("Practice", entry["practice"], styles),
        Spacer(1, 9),
    ]
    return KeepTogether(flowables)


def add_pdf_page_number(canvas, doc) -> None:
    if doc.page <= 1:
        return
    canvas.saveState()
    canvas.setFont("PsycheBody", 7.5)
    canvas.setFillColor(HexColor("#685F6D"))
    canvas.drawCentredString(3 * RL_INCH, 0.34 * RL_INCH, str(doc.page))
    canvas.restoreState()


def build_pdf(entries: list[dict]) -> None:
    register_pdf_fonts()
    styles = pdf_styles()
    page_size = (6 * RL_INCH, 9 * RL_INCH)
    pdf = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=page_size,
        leftMargin=0.58 * RL_INCH,
        rightMargin=0.58 * RL_INCH,
        topMargin=0.58 * RL_INCH,
        bottomMargin=0.58 * RL_INCH,
        title="The Psyche Symbols",
        author="Cult of Psyche",
        subject="360 original degree symbols with Cult of Psyche interpretations",
    )
    story = []
    story.extend(
        [
            Spacer(1, 0.35 * RL_INCH),
            RLImage(str(COVER_PNG), width=4.2 * RL_INCH, height=6.3 * RL_INCH),
            PageBreak(),
            Paragraph("How to Read the Psyche Symbols", styles["h1"]),
        ]
    )
    for text in [
        "The Psyche Symbols are an original 360-degree image system built for reflective astrology, creative work, journaling, and conversation. The zodiac provides the sequence. Each degree provides a scene. The reader provides attention.",
        "A symbol can reveal a tension without issuing a command. It may show what you desire, fear, repeat, protect, or refuse to name. It cannot remove your judgment. Treat every interpretation as a doorway you may enter, question, revise, or leave.",
        "The reading anatomy has five parts. Interpretation studies the whole scene. Signal names the image's living capacity. Shadow names the distortion that may travel beside that capacity. Mirror asks a question that returns authority to the reader. Practice brings the symbol into an observable act.",
        "The Cult of Psyche holds truth and mystery together. Read the symbol as symbol. Check facts as facts. Let intuition propose; let evidence correct; let consent govern what happens next.",
    ]:
        story.append(Paragraph(escape(text), styles["front"]))
    story.extend([Spacer(1, 10), Paragraph("The Twelve Chambers", styles["h1"])])

    table_data = [
        [
            Paragraph('<font color="#FFFFFF"><b>RANGE</b></font>', styles["table"]),
            Paragraph('<font color="#FFFFFF"><b>SIGN</b></font>', styles["table"]),
            Paragraph('<font color="#FFFFFF"><b>CHAMBER</b></font>', styles["table"]),
        ]
    ]
    for sign_idx, sign in enumerate(SIGNS):
        start = sign_idx * 30 + 1
        table_data.append(
            [
                Paragraph(f"{start}-{start + 29}", styles["table"]),
                Paragraph(escape(sign), styles["table_bold"]),
                Paragraph(escape(SIGN_PROFILES[sign]["chapter"]), styles["table"]),
            ]
        )
    table = Table(table_data, colWidths=[0.7 * RL_INCH, 1.05 * RL_INCH, 3.0 * RL_INCH], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), HexColor("#2D1547")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#FCF8EE"), colors.white]),
                ("LINEBELOW", (0, 0), (-1, -2), 0.25, HexColor("#D7CFBF")),
            ]
        )
    )
    story.extend([table, PageBreak(), Paragraph("Editorial Boundary", styles["h1"])])
    for text in [
        "This is a project-owned original work. It does not reproduce historical Sabian symbol wording and should not be presented as an authorized historical Sabian edition.",
        "The public CultCodex archive informed the book's vocabulary of signal, shadow, symbolic interpretation, consciousness, sovereignty, and living memory. CultCodex marks Psychenomicon material as mythic or symbolic interpretation rather than factual reporting. This book follows that boundary and excludes claims about real people.",
        "Public reference consulted September 9, 2026: cultcodex.me and cultcodex.me/about/methodology",
    ]:
        story.append(Paragraph(escape(text), styles["front"]))
    story.append(PageBreak())

    for sign in SIGNS:
        profile = SIGN_PROFILES[sign]
        story.extend(
            [
                Spacer(1, 1.25 * RL_INCH),
                Paragraph(escape(sign), styles["h1_center"]),
                Paragraph(escape(profile["chapter"]), styles["h2_center"]),
                RLImage(str(FRONTISPIECES_DIR / f"{SIGNS.index(sign) + 1:02d}-{sign}.png"), width=2.65 * RL_INCH, height=3.53 * RL_INCH),
                Paragraph(escape(" ".join(profile["movement"])), styles["chapter_intro"]),
            ]
        )
        phase_data = [
            [
                Paragraph('<font color="#FFFFFF"><b>DEGREES</b></font>', styles["table"]),
                Paragraph('<font color="#FFFFFF"><b>GATE</b></font>', styles["table"]),
                Paragraph('<font color="#FFFFFF"><b>WORK</b></font>', styles["table"]),
            ]
        ]
        for degree_range, phase, _, practice in PHASES:
            phase_data.append(
                [
                    Paragraph(f"{min(degree_range)}-{max(degree_range)}", styles["table"]),
                    Paragraph(escape(phase), styles["table_bold"]),
                    Paragraph(escape(practice), styles["table"]),
                ]
            )
        phase_table = Table(phase_data, colWidths=[0.6 * RL_INCH, 0.85 * RL_INCH, 3.3 * RL_INCH], repeatRows=1)
        phase_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), HexColor("#2D1547")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#FCF8EE"), colors.white]),
                    ("LINEBELOW", (0, 0), (-1, -2), 0.25, HexColor("#D7CFBF")),
                ]
            )
        )
        story.extend([phase_table, PageBreak()])
        sign_entries = [entry for entry in entries if entry["sign"] == sign]
        for idx, entry in enumerate(sign_entries):
            story.append(pdf_symbol_block(entry, styles))
            if idx % 2 == 1 and idx != len(sign_entries) - 1:
                story.append(PageBreak())
        story.append(PageBreak())

    story.extend(
        [
            Spacer(1, 1.15 * RL_INCH),
            Paragraph("The Signal Continues", styles["h1_center"]),
        ]
    )
    for text in [
        "The final degree leaves a microphone waiting at the edge of dawn. The book closes there because a living symbol completes its work in another mind.",
        "Return to any image when its detail begins moving again. Read slowly. Keep what clarifies. Question what flatters. Put down whatever tries to take your judgment from you.",
        "These pages belong to fellow Travelers, each carrying a Lantern far enough to see the next honest step.",
        "The dark microphone waits for the one still arriving.",
    ]:
        story.append(Paragraph(escape(text), styles["closing"]))
    story.extend([Spacer(1, 12), Paragraph("THE PSYCHE SYMBOLS", styles["cover_small"])])
    pdf.build(story, onFirstPage=add_pdf_page_number, onLaterPages=add_pdf_page_number)


def build(entries: list[dict]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_markdown(entries)

    doc = Document()
    configure_document(doc)
    add_footer(doc.sections[0])
    add_cover(doc)
    add_front_matter(doc)
    for sign in SIGNS:
        add_sign_chapter(doc, sign, [entry for entry in entries if entry["sign"] == sign])
    add_closing(doc)

    props = doc.core_properties
    props.title = "The Psyche Symbols"
    props.subject = "360 original degree symbols with Cult of Psyche interpretations"
    props.author = "Cult of Psyche"
    props.keywords = "Psyche Symbols, astrology, symbolic interpretation, Cult of Psyche"
    props.comments = "Original 2026 edition"
    doc.save(DOCX_PATH)
    build_pdf(entries)


def content_qa(entries: list[dict]) -> dict:
    expected = list(range(1, 361))
    indices = [e["global_index"] for e in entries]
    headings = [f"{e['sign']} {e['degree']} {e['title']}" for e in entries]
    sections = [e["reading"] + e["signal"] + e["shadow"] + e["mirror"] + e["practice"] for e in entries]
    source_images = [e["image"] for e in entries]
    titles = [e["title"] for e in entries]
    counts = {sign: sum(1 for e in entries if e["sign"] == sign) for sign in SIGNS}
    banned = re.compile(
        r"\b(delve|foster|leverage|utilize|facilitate|empower|streamline|robust|cutting-edge|paradigm shift|game changer|tapestry|realm|beacon|multifaceted|meticulous|intricate|paramount|transformative|elevate|embark|supercharge|harness|ever-evolving|TODO|TBD|TK)\b",
        re.IGNORECASE,
    )
    full = "\n".join(sections + source_images)
    field_uniqueness = {
        field: len({entry[field] for entry in entries})
        for field in ["reading", "signal", "shadow", "mirror", "practice"]
    }
    interpretive_sentences = []
    for entry in entries:
        for field in ["reading", "signal", "shadow", "mirror", "practice"]:
            interpretive_sentences.extend(split_sentences(entry[field]))
    repeated_sentences = {
        sentence: count
        for sentence, count in Counter(interpretive_sentences).items()
        if count > 1 and len(sentence.split()) >= 5
    }
    stock_pattern = re.compile(
        r"\b(read the scene closely|the image is exact|opens with a concrete fact|"
        r"the signal is|the shadow appears as|this is the [a-z]+ gate|"
        r"it'?s worth noting|it'?s important to note|at the end of the day|"
        r"at its core|in conclusion|ultimately|overall)\b",
        re.IGNORECASE,
    )
    stock_phrase_hits = Counter(match.group(0).lower() for match in stock_pattern.finditer(full))
    qa = {
        "entries": len(entries),
        "sequence_ok": indices == expected,
        "sign_counts": counts,
        "all_sign_counts_30": all(value == 30 for value in counts.values()),
        "unique_headings": len(set(headings)),
        "unique_titles": len(set(titles)),
        "unique_source_images": len(set(source_images)),
        "unique_interpretive_sections": len(set(sections)),
        "field_uniqueness": field_uniqueness,
        "repeated_interpretive_sentences": repeated_sentences,
        "stock_phrase_hits": dict(stock_phrase_hits),
        "source_scene_word_range": [
            min(len(image.split()) for image in source_images),
            max(len(image.split()) for image in source_images),
        ],
        "source_scenes_are_sentences": all(image[:1].isupper() and image.endswith((".", "?", "!")) for image in source_images),
        "entries_with_all_fields": sum(
            1 for e in entries if all(e[field].strip() for field in ["reading", "signal", "shadow", "mirror", "practice"])
        ),
        "banned_word_hits": sorted(set(match.group(0).lower() for match in banned.finditer(full))),
        "markdown_path": str(MARKDOWN_PATH),
        "docx_path": str(DOCX_PATH),
        "pdf_path": str(PDF_PATH),
    }
    QA_PATH.write_text(json.dumps(qa, indent=2), encoding="utf-8")
    return qa


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    symbols = parse_symbols()
    entries = [interpret(symbol) for symbol in symbols]
    qa = content_qa(entries)
    if not (
        qa["entries"] == 360
        and qa["sequence_ok"]
        and qa["all_sign_counts_30"]
        and qa["unique_headings"] == 360
        and qa["unique_titles"] == 360
        and qa["unique_source_images"] == 360
        and qa["unique_interpretive_sections"] == 360
        and all(count == 360 for count in qa["field_uniqueness"].values())
        and not qa["repeated_interpretive_sentences"]
        and not qa["stock_phrase_hits"]
        and qa["source_scenes_are_sentences"]
        and qa["entries_with_all_fields"] == 360
        and not qa["banned_word_hits"]
    ):
        raise SystemExit(json.dumps(qa, indent=2))
    build(entries)
    print(json.dumps(qa, indent=2))


if __name__ == "__main__":
    main()
